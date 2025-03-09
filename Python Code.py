import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import os
from docx import Document
import openpyxl

# Function to read the Word document
def read_word_file(word_file_path):
    doc = Document(word_file_path)
    word_text = ""
    for para in doc.paragraphs:
        word_text += para.text + "\n"
    return word_text

# Function to read email addresses and names from Excel file
def read_email_addresses_from_excel(excel_file_path):
    wb = openpyxl.load_workbook(excel_file_path)
    sheet = wb.active
    email_list = []

    # Assuming email addresses are in column A and names in column B starting from row 2
    for row in range(2, sheet.max_row + 1):
        email = sheet.cell(row=row, column=1).value
        name = sheet.cell(row=row, column=2).value  # Read names from column B
        if email and name:
            email_list.append((email, name))  # Store both email and name as a tuple

    return email_list

# Function to send an email
def send_email(recipient, subject, body, attachment_path):
    sender_email = "xyz@gmail.com"  
    sender_password = "Passowrd"  
    smtp_server = "smtp.gmail.com"
    smtp_port = 587

    # Create a MIME object for the email
    msg = MIMEMultipart()
    msg['From'] = sender_email
    msg['To'] = recipient
    msg['Subject'] = subject

    # Add the email body
    msg.attach(MIMEText(body, 'plain'))

    # Attach the file if the path is valid
    if attachment_path and os.path.exists(attachment_path):
        with open(attachment_path, "rb") as attachment:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f"attachment; filename={os.path.basename(attachment_path)}")
            msg.attach(part)

    # Send the email using SMTP
    with smtplib.SMTP(smtp_server, smtp_port) as server:
        server.starttls()  # Secure the connection
        server.login(sender_email, sender_password)  # Login to the Gmail account
        server.sendmail(sender_email, recipient, msg.as_string())  # Send the email
        print(f"Email sent to {recipient}")

# Path to the Word document and attachment
word_file_path = ""  # Change this path
attachment_path = ""  # Change this path

# Read the body from the Word document
email_body = read_word_file(word_file_path)

# Excel file containing email addresses and names
excel_file_path = "/Users/prithvichauhan/Desktop/Book1.xlsx"  # Change this path to your Excel file

# Read email addresses and names from the Excel file
email_list = read_email_addresses_from_excel(excel_file_path)

# Email details
subject = "Subject "

# Loop through the email list and send personalized emails
for recipient, name in email_list:
    personalized_body = email_body.replace("[Name]", name)  # Replace [Name] with recipient's name
    send_email(recipient, subject, personalized_body, attachment_path)

print("Emails sent successfully!")
